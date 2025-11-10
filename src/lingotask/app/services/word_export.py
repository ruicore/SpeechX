from datetime import datetime
from pathlib import Path
from typing import Iterable

from app.schemas import VocabularyItem
from docx import Document


def build_docx(*, text: str, translation: str, vocabulary: Iterable[VocabularyItem], download_dir: Path) -> Path:
    doc = Document()
    doc.add_heading('Source Text', level=2)
    doc.add_paragraph(text)
    doc.add_heading('Chinese Translation', level=2)
    doc.add_paragraph(translation)

    doc.add_heading('Professional Vocabulary', level=2)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'English'
    hdr_cells[1].text = 'Chinese'
    hdr_cells[2].text = 'Explanation'
    for v in vocabulary:
        row = table.add_row().cells
        row[0].text = v.english
        row[1].text = v.chinese
        row[2].text = v.explanation

    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    out_path = download_dir / f"translation_{ts}.docx"
    doc.save(out_path)
    return out_path
